import os
import io
import datetime
from ..template_env import templates
from fastapi import APIRouter, Request, Depends, Form
from fastapi.responses import RedirectResponse, StreamingResponse
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, Alignment

from ..database import get_db
from ..auth import get_current_user
from ..models import Transaction, TransactionLine, Account, AuditLog, FinancialNote
from .. import accounting as acc

router = APIRouter(prefix="/reports")


def _today():
    return datetime.date.today()


def _parse(d, fallback):
    if not d:
        return fallback
    return datetime.date.fromisoformat(d)


def _log_report(db, user, name, period_str):
    db.add(AuditLog(username=user.username, action=f"Generated {name} ({period_str})", module="REPORTS"))
    db.commit()


def _xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _header(ws, title, period_line):
    ws["A1"] = "RELAXIFY LOUNGE"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = title
    ws["A2"].font = Font(bold=True, size=12)
    ws["A3"] = period_line
    ws["A3"].font = Font(italic=True, size=10)


# ---------------- TRIAL BALANCE ----------------

@router.get("/trial-balance")
def trial_balance_view(request: Request, db: Session = Depends(get_db), as_of: str = None, export: str = None):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    as_of_date = _parse(as_of, _today())
    rows, total_debit, total_credit = acc.trial_balance(db, as_of_date)

    if export == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Trial Balance"
        _header(ws, "Trial Balance", f"As of {as_of_date.strftime('%B %d, %Y')}")
        ws.append([])
        ws.append(["Account Code", "Account Name", "Debit", "Credit"])
        for c in ws[5]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r["code"], r["name"], float(r["debit"]), float(r["credit"])])
        ws.append(["", "TOTAL", float(total_debit), float(total_credit)])
        ws[f"B{ws.max_row}"].font = Font(bold=True)
        _log_report(db, user, "Trial Balance", f"as of {as_of_date}")
        return _xlsx_response(wb, f"TrialBalance_{as_of_date}.xlsx")

    _log_report(db, user, "Trial Balance", f"as of {as_of_date}")
    return templates.TemplateResponse("reports/trial_balance.html", {
        "request": request, "user": user, "rows": rows,
        "total_debit": total_debit, "total_credit": total_credit,
        "as_of": as_of_date.isoformat(), "is_balanced": total_debit == total_credit,
    })


# ---------------- INCOME STATEMENT ----------------

@router.get("/income-statement")
def income_statement_view(request: Request, db: Session = Depends(get_db), date_from: str = None, date_to: str = None, export: str = None):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    d_from = _parse(date_from, _today().replace(day=1))
    d_to = _parse(date_to, _today())
    data = acc.income_statement(db, d_from, d_to)

    if export == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Income Statement"
        _header(ws, "Income Statement", f"{d_from.strftime('%B %d, %Y')} to {d_to.strftime('%B %d, %Y')}")
        ws.append([])
        ws.append(["Revenue"])
        for r in data["revenue"]:
            ws.append([r["code"], r["name"], float(r["amount"])])
        ws.append(["", "Total Revenue", float(data["total_revenue"])])
        ws.append([])
        ws.append(["Cost of Goods Sold"])
        for r in data["cogs"]:
            ws.append([r["code"], r["name"], float(r["amount"])])
        ws.append(["", "Gross Profit", float(data["gross_profit"])])
        ws.append([])
        ws.append(["Expenses"])
        for r in data["expenses"]:
            ws.append([r["code"], r["name"], float(r["amount"])])
        ws.append(["", "Total Expenses", float(data["total_expenses"])])
        ws.append([])
        ws.append(["", "NET INCOME / (LOSS)", float(data["net_income"])])
        ws[f"B{ws.max_row}"].font = Font(bold=True)
        _log_report(db, user, "Income Statement", f"{d_from} to {d_to}")
        return _xlsx_response(wb, f"IncomeStatement_{d_from}_to_{d_to}.xlsx")

    _log_report(db, user, "Income Statement", f"{d_from} to {d_to}")
    return templates.TemplateResponse("reports/income_statement.html", {
        "request": request, "user": user, "data": data,
        "date_from": d_from.isoformat(), "date_to": d_to.isoformat(),
    })


# ---------------- BALANCE SHEET ----------------

@router.get("/balance-sheet")
def balance_sheet_view(request: Request, db: Session = Depends(get_db), as_of: str = None, export: str = None):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    as_of_date = _parse(as_of, _today())
    data = acc.balance_sheet(db, as_of_date)

    if export == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Balance Sheet"
        _header(ws, "Balance Sheet", f"As of {as_of_date.strftime('%B %d, %Y')}")
        ws.append([])
        ws.append(["Assets"])
        for r in data["assets"]:
            ws.append([r["code"], r["name"], float(r["amount"])])
        ws.append(["", "Total Assets", float(data["total_assets"])])
        ws.append([])
        ws.append(["Liabilities"])
        for r in data["liabilities"]:
            ws.append([r["code"], r["name"], float(r["amount"])])
        ws.append(["", "Total Liabilities", float(data["total_liabilities"])])
        ws.append([])
        ws.append(["Equity"])
        for r in data["equity"]:
            ws.append([r["code"], r["name"], float(r["amount"])])
        ws.append(["", "Total Equity", float(data["total_equity"])])
        ws.append([])
        ws.append(["", "Total Liabilities + Equity", float(data["total_liabilities_and_equity"])])
        _log_report(db, user, "Balance Sheet", f"as of {as_of_date}")
        return _xlsx_response(wb, f"BalanceSheet_{as_of_date}.xlsx")

    _log_report(db, user, "Balance Sheet", f"as of {as_of_date}")
    return templates.TemplateResponse("reports/balance_sheet.html", {
        "request": request, "user": user, "data": data, "as_of": as_of_date.isoformat(),
    })


# ---------------- STATEMENT OF CHANGES IN EQUITY ----------------

@router.get("/equity-changes")
def equity_changes_view(request: Request, db: Session = Depends(get_db), date_from: str = None, date_to: str = None):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    d_from = _parse(date_from, _today().replace(day=1))
    d_to = _parse(date_to, _today())
    data = acc.statement_of_changes_in_equity(db, d_from, d_to)
    _log_report(db, user, "Statement of Changes in Equity", f"{d_from} to {d_to}")
    return templates.TemplateResponse("reports/equity_changes.html", {
        "request": request, "user": user, "data": data,
        "date_from": d_from.isoformat(), "date_to": d_to.isoformat(),
    })


# ---------------- CASH FLOWS ----------------

@router.get("/cash-flows")
def cash_flows_view(request: Request, db: Session = Depends(get_db), date_from: str = None, date_to: str = None):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    d_from = _parse(date_from, _today().replace(day=1))
    d_to = _parse(date_to, _today())
    data = acc.cash_flow_statement(db, d_from, d_to)
    _log_report(db, user, "Statement of Cash Flows", f"{d_from} to {d_to}")
    return templates.TemplateResponse("reports/cash_flows.html", {
        "request": request, "user": user, "data": data,
        "date_from": d_from.isoformat(), "date_to": d_to.isoformat(),
    })


# ---------------- NOTES ----------------

@router.get("/notes")
def notes_view(request: Request, db: Session = Depends(get_db)):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    notes = db.query(FinancialNote).order_by(FinancialNote.id).all()
    return templates.TemplateResponse("reports/notes.html", {"request": request, "user": user, "notes": notes})


@router.post("/notes/add")
def add_note(request: Request, db: Session = Depends(get_db), title: str = Form(...), content: str = Form(...)):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    db.add(FinancialNote(title=title.strip(), content=content.strip(), updated_by=user.username))
    db.add(AuditLog(username=user.username, action=f"Added financial note: {title}", module="REPORTS"))
    db.commit()
    return RedirectResponse("/reports/notes", status_code=303)


@router.post("/notes/{note_id}/edit")
def edit_note(note_id: int, request: Request, db: Session = Depends(get_db), title: str = Form(...), content: str = Form(...)):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    note = db.query(FinancialNote).filter(FinancialNote.id == note_id).first()
    if note:
        note.title = title.strip()
        note.content = content.strip()
        note.updated_by = user.username
        note.updated_at = datetime.datetime.utcnow()
        db.add(AuditLog(username=user.username, action=f"Edited financial note: {note.title}", module="REPORTS"))
        db.commit()
    return RedirectResponse("/reports/notes", status_code=303)


# ---------------- DAILY TRANSACTIONS ----------------

@router.get("/daily")
def daily_view(request: Request, db: Session = Depends(get_db), date_from: str = None, date_to: str = None, export: str = None):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)
    d_from = _parse(date_from, _today())
    d_to = _parse(date_to, _today())

    lines = (
        db.query(TransactionLine)
        .join(Transaction, Transaction.id == TransactionLine.transaction_id)
        .filter(Transaction.transaction_date >= d_from, Transaction.transaction_date <= d_to)
        .order_by(Transaction.transaction_date, Transaction.id)
        .all()
    )

    if export == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Daily Transactions"
        _header(ws, "Daily Transaction Report", f"{d_from.strftime('%B %d, %Y')} to {d_to.strftime('%B %d, %Y')}")
        ws.append([])
        ws.append(["Date", "Reference No.", "Type", "Account Code", "Account Name", "Debit", "Credit", "Remarks", "Entered By", "Status", "Recorded At"])
        for c in ws[5]:
            c.font = Font(bold=True)
        for l in lines:
            t = l.transaction
            ws.append([
                t.transaction_date.strftime("%m/%d/%Y"), t.reference_number, t.transaction_type,
                l.account.code, l.account.name, float(l.debit), float(l.credit),
                t.remarks or "", t.created_by, t.status, (t.created_at + datetime.timedelta(hours=8)).strftime("%m/%d/%Y %H:%M"),
            ])
        _log_report(db, user, "Daily Transaction Report", f"{d_from} to {d_to}")
        return _xlsx_response(wb, f"DailyTransactions_{d_from}_to_{d_to}.xlsx")

    _log_report(db, user, "Daily Transaction Report", f"{d_from} to {d_to}")
    return templates.TemplateResponse("reports/daily.html", {
        "request": request, "user": user, "lines": lines,
        "date_from": d_from.isoformat(), "date_to": d_to.isoformat(),
    })


# ---------------- ACCOUNT LEDGER (search any account, see its history) ----------------

@router.get("/account-ledger")
def account_ledger_view(request: Request, db: Session = Depends(get_db), account_id: int = None, date_from: str = None, date_to: str = None):
    user = get_current_user(request, db)
    if not user:
        return RedirectResponse("/login", status_code=303)

    all_accounts = db.query(Account).order_by(Account.code).all()
    parent_ids = {a.parent_id for a in all_accounts if a.parent_id}
    d_from = _parse(date_from, _today().replace(day=1))
    d_to = _parse(date_to, _today())

    data = None
    selected_account = None
    if account_id:
        selected_account = db.query(Account).filter(Account.id == account_id).first()
        if selected_account:
            data = acc.account_ledger(db, selected_account, d_from, d_to)
            _log_report(db, user, f"Account Ledger for {selected_account.code} - {selected_account.name}", f"{d_from} to {d_to}")

    return templates.TemplateResponse("reports/account_ledger.html", {
        "request": request, "user": user, "all_accounts": all_accounts, "parent_ids": parent_ids,
        "selected_account": selected_account, "data": data,
        "date_from": d_from.isoformat(), "date_to": d_to.isoformat(),
    })
